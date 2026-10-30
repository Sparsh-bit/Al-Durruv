"""Report rendering: PDF (client review) and XLSX (data appendix)."""

import io
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.ai.grounding import fmt_usd
from app.exports import safe_cell

NAVY = colors.HexColor("#0F2A4A")
MUTED = colors.HexColor("#5B6B7F")
LABEL_COLORS = {
    "verified": "#0E7A4F",
    "client_provided": "#7A4FB5",
    "calculated": "#0F6E84",
    "estimate": "#B26A00",
    "inference": "#3756C8",
    "unavailable": "#6B7280",
}
LABEL_TEXT = {k: k.replace("_", " ").upper() for k in LABEL_COLORS}


def client_statements(section: dict) -> list[dict]:
    """Statements that failed validation are never exported to the client."""
    return [s for s in section["statements"] if s["label"] != "unverified"]


def _styles():
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], textColor=NAVY, fontSize=20, leading=24, alignment=TA_LEFT),
        "meta": ParagraphStyle("m", parent=ss["Normal"], textColor=MUTED, fontSize=8.5, leading=11),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], textColor=NAVY, fontSize=13, spaceBefore=10, spaceAfter=4),
        "body": ParagraphStyle("b", parent=ss["Normal"], fontSize=9.5, leading=13.5, spaceAfter=3),
        "small": ParagraphStyle("s", parent=ss["Normal"], fontSize=7.5, leading=9.5, textColor=MUTED),
        "cell": ParagraphStyle("c", parent=ss["Normal"], fontSize=8, leading=10),
    }


def _table(rows: list[list], widths: list[float], st) -> Table:
    data = [[Paragraph(escape(str(c)), st["cell"]) for c in r] for r in rows]
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF6")),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DCE6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return t


def render_pdf(report: dict) -> bytes:
    c = report["content"]
    st = _styles()
    buf = io.BytesIO()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(18 * mm, 10 * mm, f"{c['title']} · generated {c['generated_at'][:10]} · page {doc.page}")
        canvas.drawRightString(
            192 * mm, 10 * mm, "Country-level statistics: UN Comtrade. AI text grounded and validated."
        )
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=c["title"],
        author="Trade Intelligence Platform",
    )
    els = [
        Paragraph(escape(c["title"]), st["title"]),
        Paragraph(escape(f"Product: {c['product']} · HS {c['hs_code']} — {c['hs_description']}"), st["meta"]),
        Paragraph(
            escape(
                f"Generated {c['generated_at'][:16].replace('T', ' ')} UTC · "
                f"{c['validation']['exported']} statements exported, "
                f"{c['validation']['withheld']} withheld by the validator"
            ),
            st["meta"],
        ),
        Spacer(1, 4),
        Paragraph(
            " ".join(
                f'<font color="{LABEL_COLORS[k]}"><b>{v}</b></font> ' + d
                for k, v, d in [
                    ("verified", "VERIFIED", "stated in cited data;"),
                    ("client_provided", "CLIENT PROVIDED", "from your own information;"),
                    ("calculated", "CALCULATED", "computed from sourced data;"),
                    ("estimate", "ESTIMATE", "labelled estimate;"),
                    ("inference", "INFERENCE", "analyst reasoning;"),
                    ("unavailable", "UNAVAILABLE", "data not available."),
                ]
            ),
            st["small"],
        ),
        Spacer(1, 6),
    ]

    for sec in c["sections"]:
        block = [Paragraph(escape(sec["title"]), st["h2"])]
        for s in client_statements(sec):
            refs = (
                f' <font color="#5B6B7F" size="7.5">[{", ".join(s["evidence_ids"])}]</font>'
                if s["evidence_ids"]
                else ""
            )
            block.append(
                Paragraph(
                    f'<font color="{LABEL_COLORS[s["label"]]}" size="7"><b>{LABEL_TEXT[s["label"]]}'
                    f"</b></font>&nbsp; {escape(s['text'])}{refs}",
                    st["body"],
                )
            )
        els.append(KeepTogether(block[:3]))
        els.extend(block[3:])

    if c.get("ranking"):
        els += [
            Paragraph("Appendix A — Market opportunity ranking (deterministic)", st["h2"]),
            Paragraph(
                "Score 0–100 from import size, growth, supplier openness, price level and data "
                "completeness. Weights: "
                + ", ".join(f"{k.replace('_', ' ')} {v:g}" for k, v in c["weights"].items())
                + ".",
                st["small"],
            ),
            _table(
                [["#", "Market", "Score", "Imports (latest)", "Year"]]
                + [
                    [r["rank"], r["name"], f"{r['score']:.1f}", fmt_usd(r["latest_value_usd"]), r["latest_year"]]
                    for r in c["ranking"]
                ],
                [10 * mm, 60 * mm, 20 * mm, 45 * mm, 20 * mm],
                st,
            ),
        ]
    if c.get("companies"):
        els += [
            Paragraph("Appendix B — Companies identified (public web sources)", st["h2"]),
            Paragraph(
                "Candidates are AI-discovered from public pages and must be qualified before outreach. "
                "No purchase volumes or budgets are implied.",
                st["small"],
            ),
            _table(
                [["Company", "Country", "Roles", "Status", "Source"]]
                + [
                    [
                        co["name"],
                        co["country_iso3"],
                        ", ".join(co["roles"]),
                        co["review_status"],
                        co["evidence"][0]["source"]["url"] if co["evidence"] else "—",
                    ]
                    for co in c["companies"]
                ],
                [42 * mm, 16 * mm, 32 * mm, 18 * mm, 66 * mm],
                st,
            ),
        ]
    cited = {i for sec in c["sections"] for s in client_statements(sec) for i in s["evidence_ids"]}
    ev_rows = [e for e in c["evidence"] if e["id"] in cited]
    if ev_rows:
        els += [
            Paragraph("Appendix C — Evidence and sources", st["h2"]),
            _table(
                [["ID", "Evidence", "Source"]]
                + [
                    [
                        e["id"],
                        e["text"],
                        (e.get("source") or {}).get("url", "Computed from the data above")
                        + (
                            f" (retrieved {e['source']['retrieved_at'][:10]})"
                            if e.get("source") and e["source"].get("retrieved_at")
                            else ""
                        ),
                    ]
                    for e in ev_rows
                ],
                [12 * mm, 92 * mm, 70 * mm],
                st,
            ),
        ]
    doc.build(els, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def render_xlsx(report: dict) -> bytes:
    c = report["content"]
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F2A4A")

    def sheet(ws, header: list[str], rows: list[list]):
        ws.append([safe_cell(h) for h in header])
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for r in rows:
            ws.append([safe_cell(x) for x in r])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(x.value or "")) for x in col)))

    ws = wb.active
    ws.title = "Brief"
    sheet(
        ws,
        ["Section", "Label", "Statement", "Evidence"],
        [
            [sec["title"], s["label"], s["text"], ", ".join(s["evidence_ids"])]
            for sec in c["sections"]
            for s in client_statements(sec)
        ],
    )
    sheet(
        wb.create_sheet("Ranking"),
        ["Rank", "Market", "ISO3", "Score", "Imports USD", "Year"],
        [
            [r["rank"], r["name"], r["iso3"], r["score"], r["latest_value_usd"], r["latest_year"]]
            for r in c.get("ranking", [])
        ],
    )
    sheet(
        wb.create_sheet("Companies"),
        ["Company", "Country", "Roles", "Status", "Website", "Source URL", "Retrieved"],
        [
            [
                co["name"],
                co["country_iso3"],
                ", ".join(co["roles"]),
                co["review_status"],
                co["website"],
                co["evidence"][0]["source"]["url"] if co["evidence"] else "",
                co["evidence"][0]["source"]["retrieved_at"][:10] if co["evidence"] else "",
            ]
            for co in c.get("companies", [])
        ],
    )
    sheet(
        wb.create_sheet("Evidence"),
        ["ID", "Kind", "Evidence", "Source URL", "Retrieved"],
        [
            [
                e["id"],
                e["kind"],
                e["text"],
                (e.get("source") or {}).get("url", ""),
                ((e.get("source") or {}).get("retrieved_at") or "")[:10],
            ]
            for e in c["evidence"]
        ],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(report: dict, ext: str) -> str:
    stamp = datetime.fromisoformat(report["content"]["generated_at"]).strftime("%Y%m%d")
    return f"market-brief-HS{report['content']['hs_code']}-{stamp}.{ext}"


# --- company account brief (English or Arabic) ------------------------------------------------------

AR_FONT = "NotoNaskhArabic"
_FONT_PATH = Path(__file__).parent.parent / "static" / "fonts" / "NotoNaskhArabic.ttf"


def _arabic_lines(text: str, size: float, width: float) -> list[str]:
    """Wrap in logical order, then shape and reorder each line, so multi-line Arabic reads top-to-bottom."""
    import arabic_reshaper
    from bidi import get_display
    from reportlab.pdfbase.pdfmetrics import stringWidth

    lines, cur = [], ""
    for word in text.split():
        trial = f"{cur} {word}".strip()
        if cur and stringWidth(arabic_reshaper.reshape(trial), AR_FONT, size) > width:
            lines.append(cur)
            cur = word
        else:
            cur = trial
    if cur:
        lines.append(cur)
    return [get_display(arabic_reshaper.reshape(line)) for line in lines]


def _src_label(src: dict | None) -> str:
    src = src or {}
    return src.get("url") or src.get("title", "")


def render_company_brief_pdf(content: dict) -> bytes:
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    rtl = content.get("language") == "ar"
    if rtl and AR_FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(AR_FONT, str(_FONT_PATH)))
    st = _styles()
    buf = io.BytesIO()
    width = A4[0] - 36 * mm
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=f"Account brief: {content['company']}",
        author="Trade Intelligence Platform",
    )

    def para(text: str, style: str) -> list:
        if not rtl:
            return [Paragraph(escape(text), st[style])]
        base = st[style]
        s = ParagraphStyle(
            f"ar-{style}", parent=base, fontName=AR_FONT, alignment=TA_RIGHT, leading=base.fontSize * 1.6
        )
        return [Paragraph(escape(line), s) for line in _arabic_lines(text, base.fontSize, width)]

    els = [
        Paragraph(escape(content["company"]), st["title"]),
        Paragraph(
            escape(
                f"Generated {content['generated_at'][:16].replace('T', ' ')} UTC · "
                f"language {content.get('language', 'en')}"
            ),
            st["meta"],
        ),
        Spacer(1, 6),
    ]
    for sec in content["sections"]:
        els += para(sec["title"], "h2")
        for s in sec["statements"]:
            if s["label"] == "unverified":
                continue
            refs = f" [{s['label']}; {', '.join(s['evidence_ids'])}]" if s["evidence_ids"] else f" [{s['label']}]"
            els += para(s["text"] + ("" if rtl else refs), "body")
            if rtl:
                els.append(Paragraph(escape(refs.strip()), st["small"]))
    email = content.get("outreach_email") or {}
    if email:
        els += para("Email" if not rtl else "البريد الإلكتروني", "h2")
        els += para(email.get("subject", ""), "body")
        for chunk in (email.get("body") or "").split("\n"):
            if chunk.strip():
                els += para(chunk, "body")
    if content.get("whatsapp_message"):
        els += para("WhatsApp", "h2") + para(content["whatsapp_message"], "body")
    els += [Spacer(1, 8), Paragraph("Evidence", st["h2"])]
    rows = [["ID", "Kind", "Fact", "Source"]] + [
        [
            e["id"],
            e["kind"],
            e["text"][:400],
            (e.get("source") or {}).get("url") or (e.get("source") or {}).get("title", ""),
        ]
        for e in content.get("evidence", [])
    ]
    els.append(_table(rows, [12 * mm, 20 * mm, 100 * mm, width - 132 * mm], st))
    doc.build(els)
    return buf.getvalue()
