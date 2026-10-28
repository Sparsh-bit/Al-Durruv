import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.adapters import comtrade
from app.ai import claude
from app.main import app
from app.models import ApiCache, Product, SourceRecord, TradeObservation
from app.services import research, trade

# ---------- auth & access control ----------


def test_login_and_protection(viewer):
    anon = TestClient(app)
    assert anon.post("/api/auth/login", data={"username": "admin@example.com", "password": "wrong"}).status_code == 401
    assert anon.get("/api/products").status_code == 401
    assert viewer.get("/api/products").status_code == 200
    assert viewer.post("/api/products", json={"sku": "A", "name": "A", "hs_code": "190531"}).status_code == 403
    assert viewer.get("/api/admin/runs").status_code == 403
    assert viewer.get("/api/admin/sources").status_code == 200  # integration health has no secrets


def test_admin_cannot_demote_self(admin):
    me = admin.get("/api/auth/me").json()
    assert admin.patch(f"/api/auth/users/{me['id']}", json={"role": "viewer"}).status_code == 400


# ---------- catalogue ----------


def test_product_crud_validates_hs(analyst):
    assert analyst.post("/api/products", json={"sku": "X1", "name": "X", "hs_code": "999999"}).status_code == 422
    r = analyst.post("/api/products", json={"sku": "X1", "name": "Biscuits", "hs_code": "1905.31"})
    assert r.status_code == 201 and r.json()["hs_code"] == "190531"
    assert analyst.post("/api/products", json={"sku": "X1", "name": "Dup", "hs_code": "190531"}).status_code == 409


def test_csv_import_reports_row_errors_and_restores_leading_zero(analyst, db):
    csv = b"SKU,Name,HS Code\nC1,Coffee,90111\nB1,Biscuits,190531.0\nZ1,Bad,123\n,NoSku,190531\n"
    r = analyst.post("/api/products/import", files={"file": ("p.csv", csv, "text/csv")})
    body = r.json()
    assert r.status_code == 200 and body["created"] == 2 and len(body["errors"]) == 2
    assert {e["line"] for e in body["errors"]} == {4, 5}
    assert db.scalar(select(Product.hs_code).where(Product.sku == "C1")) == "090111"
    # re-import updates instead of duplicating
    r = analyst.post(
        "/api/products/import", files={"file": ("p.csv", b"sku,name,hs_code\nC1,Coffee beans,090111\n", "text/csv")}
    )
    assert r.json() == {"created": 0, "updated": 1, "errors": []}


def test_import_rejects_missing_columns_and_bad_type(analyst):
    r = analyst.post("/api/products/import", files={"file": ("p.csv", b"sku\nA\n", "text/csv")})
    assert r.status_code == 422 and "name" in r.json()["detail"]
    assert analyst.post("/api/products/import", files={"file": ("p.txt", b"x", "text/plain")}).status_code == 415


# ---------- comtrade adapter: cache, retries, provenance ----------

ROWS = [
    {
        "refYear": 2023,
        "reporterCode": 826,
        "reporterISO": "GBR",
        "reporterDesc": "United Kingdom",
        "partnerCode": 0,
        "partnerISO": "W00",
        "partnerDesc": "World",
        "flowCode": "M",
        "primaryValue": 500e6,
        "netWgt": 1e8,
        "qty": 1e8,
        "qtyUnitAbbr": "kg",
    },
]


def test_fetch_caches_and_records_provenance(db):
    calls = []

    def fake(params, key):
        calls.append(params)
        return ROWS

    rows1, src1 = comtrade.fetch(db, period=2023, cmd="190531", flow="M", caller=fake)
    rows2, src2 = comtrade.fetch(db, period=2023, cmd="190531", flow="M", caller=fake)
    assert len(calls) == 1 and rows1 == rows2 and src1.id == src2.id
    assert src1.url.startswith("https://comtradeapi.un.org/public/v1/preview/C/A/HS?")
    assert "subscription" not in src1.url and src1.reporting_period == "2023" and src1.licence
    assert db.scalar(select(ApiCache).where(ApiCache.adapter == "un_comtrade")) is not None


def test_fetch_retries_then_raises(db, monkeypatch):
    monkeypatch.setattr(comtrade.time, "sleep", lambda s: None)
    attempts = []

    def flaky(params, key):
        attempts.append(1)
        if len(attempts) < 3:
            raise comtrade.ComtradeError("429 rate limit")
        return ROWS

    rows, _ = comtrade.fetch(db, period=2022, cmd="190531", flow="M", caller=flaky)
    assert len(attempts) == 3 and rows == ROWS

    def down(params, key):
        raise comtrade.ComtradeError("503")

    with pytest.raises(comtrade.ComtradeError, match="after 3 attempts"):
        comtrade.fetch(db, period=2021, cmd="190531", flow="M", caller=down)


def test_client_error_output_becomes_typed_error(monkeypatch):
    def prints_and_returns_none(**kw):
        print('{"statusCode": 429, "message": "Rate limit is exceeded"}')
        return None

    monkeypatch.setattr(comtrade.comtradeapicall, "previewFinalData", prints_and_returns_none)
    with pytest.raises(comtrade.ComtradeError, match="Rate limit"):
        comtrade._call({"period": "2023", "cmdCode": "190531", "flowCode": "M"}, "")


def _load(db):
    """Synthetic but realistic observations for two markets, with supplier detail."""
    rows = []
    for year, gbr, deu in [(2020, 400e6, 600e6), (2021, 420e6, 610e6), (2022, 460e6, 640e6), (2023, 500e6, 650e6)]:
        for code, iso, name, v in [(826, "GBR", "United Kingdom", gbr), (276, "DEU", "Germany", deu)]:
            rows.append(
                {
                    **ROWS[0],
                    "refYear": year,
                    "reporterCode": code,
                    "reporterISO": iso,
                    "reporterDesc": name,
                    "primaryValue": v,
                    "netWgt": v / 4,
                }
            )
    for pcode, piso, pname, v in [(484, "MEX", "Mexico", 300e6), (528, "NLD", "Netherlands", 200e6)]:
        rows.append(
            {
                **ROWS[0],
                "reporterCode": 826,
                "partnerCode": pcode,
                "partnerISO": piso,
                "partnerDesc": pname,
                "primaryValue": v,
                "netWgt": v / 4,
            }
        )
    src = SourceRecord(adapter="un_comtrade", url="https://comtradeapi.un.org/test", reporting_period="2023")
    db.add(src)
    db.flush()
    trade._store(db, rows, "190531", src.id)
    db.commit()


def test_every_observation_has_provenance_and_world_is_not_a_country(db, viewer):
    _load(db)
    obs = db.scalars(select(TradeObservation)).all()
    assert obs and all(o.source_record_id for o in obs)
    ov = viewer.get("/api/markets/190531/overview").json()
    assert ov["has_data"] and ov["reference_year"] == 2023 and ov["world_imports_usd"] == pytest.approx(1150e6)
    assert ov["top_importers"][0]["source"]["url"] == "https://comtradeapi.un.org/test"
    detail = viewer.get("/api/markets/190531/countries/826").json()
    assert [s["name"] for s in detail["suppliers"]] == ["Mexico", "Netherlands"]
    assert detail["suppliers"][0]["share"] == pytest.approx(0.6)
    rank = viewer.post("/api/markets/190531/ranking", json={}).json()["markets"]
    assert {m["iso3"] for m in rank} == {"GBR", "DEU"}
    assert viewer.post("/api/markets/190531/ranking", json={k: 0 for k in rank[0]["components"]}).status_code == 422


# ---------- AI: grounding enforced at the API ----------


def test_ai_endpoints_return_503_without_key(analyst):
    r = analyst.post("/api/ai/ask", json={"hs_code": "190531", "question": "Which market?"})
    assert r.status_code == 503 and "ANTHROPIC_API_KEY" in r.json()["detail"]
    assert analyst.post("/api/research", json={"hs_code": "190531", "country_m49": 826}).status_code == 503


def test_ask_validates_model_output(db, analyst, ai_on, monkeypatch):
    _load(db)

    def fake_answer(question, evidence, history):
        uk = next(e for e in evidence if "United Kingdom imported" in e.text)
        return claude.Answer(
            statements=[
                claude.Statement(
                    text="The United Kingdom imported $500.0 million.", label="verified", evidence_ids=[uk.id]
                ),
                claude.Statement(text="UK imports were $9 billion.", label="verified", evidence_ids=[uk.id]),
                claude.Statement(text="A distributor route is advisable.", label="verified", evidence_ids=["E999"]),
            ],
            follow_up_questions=["Next?"],
        ), {"model": "test", "input_tokens": 1, "output_tokens": 1}

    monkeypatch.setattr(claude, "answer_question", fake_answer)
    body = analyst.post("/api/ai/ask", json={"hs_code": "190531", "question": "Tell me about the UK"}).json()
    labels = [s["label"] for s in body["statements"]]
    assert labels == ["verified", "unverified", "inference"]
    assert {e["id"] for e in body["evidence"]} == {body["statements"][0]["evidence_ids"][0]}


def test_research_keeps_only_candidates_with_retrieved_sources():
    results = {
        "https://acme.co.uk/about": {
            "url": "https://acme.co.uk/about/",
            "title": "Acme",
            "quotes": ["Acme imports biscuits"],
        }
    }
    kept, dropped = research.validate_candidates(
        [
            {"name": "Acme Foods Ltd", "roles": ["importer", "spaceship"], "source_url": "https://acme.co.uk/about/"},
            {"name": "Ghost Trading", "roles": ["importer"], "source_url": "https://invented.example/page"},
            {"name": "acme foods ltd", "roles": ["importer"], "source_url": "https://acme.co.uk/about"},
        ],
        results,
    )
    assert [k["name"] for k in kept] == ["Acme Foods Ltd"] and kept[0]["roles"] == ["importer"]
    assert {d["reason"] for d in dropped} == {"source_url was not among retrieved search results", "duplicate"}


# ---------- reports: generation, withholding, export ----------


def test_brief_withholds_unverified_and_exports(db, analyst, ai_on, monkeypatch):
    _load(db)
    pid = analyst.post("/api/products", json={"sku": "B1", "name": "Butter cookies", "hs_code": "190531"}).json()["id"]

    def fake_brief(product, hs, evidence):
        uk = next(e for e in evidence if "United Kingdom imported" in e.text)
        return claude.Brief(
            sections=[
                claude.BriefSection(
                    title="Executive summary",
                    statements=[
                        claude.Statement(
                            text="The United Kingdom imported $500.0 million.", label="verified", evidence_ids=[uk.id]
                        ),
                        claude.Statement(text="The UK buys 7 billion units.", label="inference", evidence_ids=[]),
                    ],
                )
            ]
        ), {"model": "test-model", "input_tokens": 1, "output_tokens": 1}

    monkeypatch.setattr(claude, "write_brief", fake_brief)
    r = analyst.post("/api/reports", json={"product_id": pid})
    assert r.status_code == 201
    rep = r.json()
    assert rep["validation"] == {"statements": 2, "exported": 1, "withheld": 1}

    pdf = analyst.get(f"/api/reports/{rep['id']}/export?format=pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    assert "attachment" in pdf.headers["content-disposition"]

    xlsx = analyst.get(f"/api/reports/{rep['id']}/export?format=xlsx")
    wb = load_workbook(io.BytesIO(xlsx.content))
    statements = [row[2] for row in wb["Brief"].iter_rows(min_row=2, values_only=True)]
    assert statements == ["The United Kingdom imported $500.0 million."]  # unverified never exported
    assert analyst.get(f"/api/reports/{rep['id']}/export?format=doc").status_code == 400


def test_brief_requires_trade_data(analyst, ai_on):
    pid = analyst.post("/api/products", json={"sku": "C1", "name": "Coffee", "hs_code": "090111"}).json()["id"]
    assert analyst.post("/api/reports", json={"product_id": pid}).status_code == 409


def test_logistics_quotes_are_labelled_client_provided(analyst, viewer):
    r = analyst.post(
        "/api/logistics",
        json={
            "origin": "Jebel Ali",
            "destination": "Felixstowe",
            "destination_iso3": "gbr",
            "mode": "sea",
            "cost_usd": 2400,
            "transit_days": 24,
            "source_note": "Forwarder quote ref 1",
        },
    )
    assert r.status_code == 201 and r.json()["provenance"] == "client_provided"
    data = viewer.get("/api/logistics?destination_iso3=GBR").json()
    assert len(data["quotes"]) == 1 and {p["status"] for p in data["providers"]} == {"not_connected"}
    assert (
        analyst.post(
            "/api/logistics",
            json={"origin": "A", "destination": "B", "destination_iso3": "GBR", "mode": "teleport", "source_note": "x"},
        ).status_code
        == 422
    )


# ---------- catalogue details, export, bulk refresh ----------


def test_import_catalogue_columns_and_hs_status_reset(analyst, db):
    csv = (
        b"sku,name,brand,category,hs_code,pack_size,net_quantity,net_unit,packaging,catalogue_page,hs_status\n"
        b'D1,Apricot Jam,Durra,Jams,190531,800 g,800,g,Glass jar,"3, 5",confirmed\n'
        b"D2,Bad qty,Durra,Jams,190531,x,abc,g,Jar,7,\n"
    )
    r = analyst.post("/api/products/import", files={"file": ("c.csv", csv, "text/csv")}).json()
    assert r["created"] == 2 and r["errors"] == []
    p = db.scalar(select(Product).where(Product.sku == "D1"))
    assert (p.brand, p.pack_size, p.net_quantity, p.packaging, p.hs_status) == (
        "Durra",
        "800 g",
        800,
        "Glass jar",
        "confirmed",
    )
    assert db.scalar(select(Product.net_quantity).where(Product.sku == "D2")) is None
    listed = next(x for x in analyst.get("/api/products").json() if x["sku"] == "D1")
    assert listed["catalogue_pages"] == [3, 5]
    # changing the HS code without confirming it again resets the status
    analyst.post(
        "/api/products/import", files={"file": ("c.csv", b"sku,name,hs_code\nD1,Apricot Jam,090111\n", "text/csv")}
    )
    db.expire_all()
    assert db.scalar(select(Product.hs_status).where(Product.sku == "D1")) == "candidate"


def test_product_validation_and_export(analyst):
    bad = analyst.post("/api/products", json={"sku": "E1", "name": "X", "hs_code": "190531", "net_unit": "lbs"})
    assert bad.status_code == 422
    ok = analyst.post(
        "/api/products",
        json={
            "sku": "E1",
            "name": "Ketchup",
            "hs_code": "190531",
            "net_quantity": 450,
            "net_unit": "g",
            "pack_size": "450 g",
            "catalogue_page": "58, 59",
        },
    )
    assert ok.status_code == 201 and ok.json()["catalogue_pages"] == [58, 59]
    exp = analyst.get("/api/products/export")
    assert exp.status_code == 200 and "attachment" in exp.headers["content-disposition"]
    assert "E1,Ketchup" in exp.content.decode("utf-8-sig")


def test_refresh_missing_queues_each_code_once(analyst, viewer, monkeypatch):
    calls = []
    monkeypatch.setattr(trade, "refresh_many", lambda jobs: calls.append(jobs))
    for sku, hs in (("A", "190531"), ("B", "190531"), ("C", "090111")):
        analyst.post("/api/products", json={"sku": sku, "name": sku, "hs_code": hs})
    first = analyst.post("/api/markets/refresh-missing").json()
    assert first == {"queued": 2, "already_active": 0} and [h for _, h in calls[0]] == ["090111", "190531"]
    assert analyst.post("/api/markets/refresh-missing").json() == {"queued": 0, "already_active": 2}
    status = viewer.get("/api/markets/status").json()
    assert status["hs_codes"] == 2 and status["loaded"] == 0 and len(status["active"]) == 2
    assert viewer.post("/api/markets/refresh-missing").status_code == 403


def test_bundled_catalogue_file_is_valid(db):
    import pandas as pd

    from app.routers.catalogue import apply_catalogue
    from app.seed import CATALOGUE

    df = pd.read_csv(CATALOGUE, dtype=str, keep_default_na=False)
    codes = sorted(set(df["hs_code"]))
    from app.models import HSCode

    db.add_all([HSCode(code=c, description=c) for c in codes if c not in ("190531", "090111")])
    db.commit()
    result = apply_catalogue(db, df)
    assert result["errors"] == [] and result["created"] == len(df) == 127
    assert df["sku"].is_unique and (df["brand"] == "Durra").all()


# ---------- quota handling and request batching ----------


def test_quota_error_is_typed_and_not_retried(db, monkeypatch):
    err = comtrade._error_from_output(
        '{"statusCode": 403, "message": "Out of call volume quota. Quota will be replenished in 00:21:48."}'
    )
    assert isinstance(err, comtrade.QuotaExceeded) and err.wait_s == 21 * 60 + 48
    assert not isinstance(comtrade._error_from_output("500 server error"), comtrade.QuotaExceeded)
    calls = []

    def over_quota(params, key):
        calls.append(1)
        raise err

    with pytest.raises(comtrade.QuotaExceeded):
        comtrade.fetch(db, period=2020, cmd="190531", flow="M", caller=over_quota)
    assert len(calls) == 1  # no pointless retries


def _start_run(db):
    from app.models import RefreshRun

    run = RefreshRun(adapter="un_comtrade", scope="HS 190531", status="queued")
    db.add(run)
    db.commit()
    return run.id


def test_refresh_pauses_on_quota_then_completes(db, monkeypatch):
    from app.models import RefreshRun

    real_fetch = comtrade.fetch
    state = {"quota_hits": 0, "periods": []}

    def fake_fetch(db_, **kw):
        if state["quota_hits"] == 0:
            state["quota_hits"] += 1
            raise comtrade.QuotaExceeded("quota", 120)
        state["periods"].append(str(kw["period"]))
        return real_fetch(db_, caller=lambda params, key: ROWS if kw.get("partner") == 0 else [], **kw)

    monkeypatch.setattr(comtrade, "fetch", fake_fetch)
    monkeypatch.setattr(trade, "_years", lambda: [2020, 2021, 2022, 2023])
    slept = []
    run_id = _start_run(db)
    trade.refresh_hs(run_id, "190531", sleep=slept.append)
    db.expire_all()
    run = db.get(RefreshRun, run_id)
    assert run.status == "succeeded", run.error
    assert sum(slept) == 180 and max(slept) <= 60  # waited out the quota (+60 s), in heartbeat-sized slices
    assert state["periods"][:4] == ["2020", "2021", "2022", "2023"]  # preview endpoint: one year per request


def test_refresh_splits_years_when_response_hits_record_cap(db, monkeypatch):
    from app.models import RefreshRun

    real_fetch = comtrade.fetch
    periods = []

    def fake_fetch(db_, **kw):
        periods.append(str(kw["period"]))
        rows = ROWS * 3 if "," in str(kw["period"]) else ROWS
        return real_fetch(db_, caller=lambda params, key: rows if kw.get("partner") == 0 else [], **kw)

    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "comtrade_subscription_key", "test-key")
    monkeypatch.setattr(comtrade, "max_records", lambda: 3)  # a 3-row response counts as truncated
    monkeypatch.setattr(comtrade, "fetch", fake_fetch)
    monkeypatch.setattr(trade, "_years", lambda: [2022, 2023])
    run_id = _start_run(db)
    trade.refresh_hs(run_id, "190531", sleep=lambda s: None)
    db.expire_all()
    assert db.get(RefreshRun, run_id).status == "succeeded"
    assert periods[:3] == ["2022,2023", "2022", "2023"]


# ---------- buyer profiles, guardrails, cost control, settings, audit ----------


def test_contact_filter_keeps_only_generic_company_mailboxes():
    site = "https://www.acme-foods.co.uk"
    assert research.clean_contact({"email": "sales@acme-foods.co.uk", "phone": "+44 20 7946 0000"}, site) == {
        "email": "sales@acme-foods.co.uk",
        "phone": "+44 20 7946 0000",
    }
    assert research.clean_contact({"email": "john.smith@acme-foods.co.uk"}, site) == {}  # personal mailbox
    assert research.clean_contact({"email": "info@gmail.com"}, site) == {}  # not the company's domain
    assert research.clean_contact({"contact_url": "javascript:alert(1)"}, site) == {}
    assert research.clean_contact("not a dict", site) == {}


def test_signals_without_retrieved_source_are_dropped():
    results = {"https://acme.co.uk/about": {"url": "https://acme.co.uk/about", "title": "Acme", "quotes": []}}
    kept, _ = research.validate_candidates(
        [
            {
                "name": "Acme",
                "roles": ["importer"],
                "source_url": "https://acme.co.uk/about",
                "private_label": "YES",
                "markets_served": ["gbr", "Ireland", "IRL"],
                "buying_signals": [
                    {"text": "Launching own-brand Levantine range", "url": "https://acme.co.uk/about"},
                    {"text": "Invented", "url": "https://elsewhere.example/x"},
                ],
            }
        ],
        results,
        ["importer"],
    )
    assert [s["text"] for s in kept[0]["buying_signals"]] == ["Launching own-brand Levantine range"]
    assert kept[0]["private_label"] == "yes" and kept[0]["markets_served"] == ["GBR", "IRL"]


def test_cost_and_budget_cap(db, analyst, ai_on):
    from app.ai import usage as ai_usage
    from app.services import settings as app_settings

    c = ai_usage.cost("claude-haiku-5-5", {"input_tokens": 1_000_000, "output_tokens": 1_000_000, "web_searches": 5})
    assert c == pytest.approx(0.10 + 0.50 + 0.05)
    app_settings.put(db, "ai_monthly_budget_usd", 0.5)
    ai_usage.record(db, None, "analyst", {"model": "claude-haiku-5-5", "input_tokens": 6_000_000})
    db.commit()
    r = analyst.post("/api/ai/ask", json={"hs_code": "190531", "question": "Which market?"})
    assert r.status_code == 429 and "budget" in r.json()["detail"]


def test_settings_validation_and_export_market_flag(db, admin, analyst):
    _load(db)
    weights = {"market_size": 30, "growth": 20, "retailer_presence": 20, "private_label_signal": 15, "evidence": 15}
    body = {
        "company_profile": {"name": "ALDurra", "brand": "Durra", "country_iso3": "syr"},
        "priority_categories": ["Pickles", " Pickles ", "Sauces"],
        "private_label_weights": weights,
        "ai_monthly_budget_usd": 25,
    }
    zero = {**body, "private_label_weights": dict.fromkeys(weights, 0)}
    assert admin.put("/api/admin/settings", json=zero).status_code == 422
    assert analyst.put("/api/admin/settings", json=body).status_code == 403
    ok = admin.put("/api/admin/settings", json=body).json()
    assert ok["priority_categories"] == ["Pickles", "Sauces"] and ok["company_profile"]["country_iso3"] == "SYR"
    assert analyst.put("/api/target-markets/XXX", json={"segment": "existing"}).status_code == 422
    assert analyst.put("/api/target-markets/gbr", json={"segment": "existing"}).status_code == 200
    markets = analyst.post("/api/markets/190531/ranking", json={}).json()["markets"]
    assert {m["iso3"]: m["existing_market"] for m in markets} == {"GBR": True, "DEU": False}
    log = admin.get("/api/admin/audit").json()
    assert any(e["entity"] == "settings" for e in log)


def _company(db, **kw):
    from app.models import Company

    c = Company(name="Acme Foods", country_iso3="GBR", roles=["importer"], hs_codes=["190531"], **kw)
    db.add(c)
    db.commit()
    return c.id


def test_review_rescores_and_is_audited(db, analyst, admin):
    _load(db)
    cid = _company(db, relationship_type="buyer", private_label="possible")
    before = analyst.patch(f"/api/companies/{cid}", json={"pipeline_stage": "contacted"}).json()
    after = analyst.patch(f"/api/companies/{cid}", json={"review_status": "verified"}).json()
    assert after["pipeline_stage"] == "contacted"
    assert after["prospect_score"] > (before["prospect_score"] or 0)
    competitor = analyst.patch(f"/api/companies/{cid}", json={"relationship_type": "competitor"}).json()
    assert competitor["prospect_score"] is None  # competitors are never scored as prospects
    assert analyst.get("/api/companies?relationship=competitor").json()[0]["id"] == cid
    assert sum(e["entity"] == "company" for e in admin.get("/api/admin/audit").json()) == 3
    assert analyst.patch(f"/api/companies/{cid}", json={"pipeline_stage": "maybe"}).status_code == 422


def test_company_brief_is_validated_and_logged(db, analyst, admin, ai_on, monkeypatch):
    from app.ai import claude as claude_mod

    _load(db)
    cid = _company(db, relationship_type="buyer")

    def fake(name, evidence):
        st = claude_mod.Statement
        return claude_mod.CompanyBrief(
            sections=[
                claude_mod.BriefSection(
                    title="Purchasing pattern and volumes",
                    statements=[
                        st(text="They buy 40 containers a month.", label="estimate", evidence_ids=[]),
                        st(text="Order volumes are unknown.", label="unavailable", evidence_ids=[evidence[-1].id]),
                    ],
                )
            ],
            outreach_email=claude_mod.Email(subject="Private label", body="We can supply 500 tonnes a month."),
        ), {"model": "claude-haiku-5-5", "input_tokens": 1000, "output_tokens": 500}

    monkeypatch.setattr(claude_mod, "write_company_brief", fake)
    b = analyst.post(f"/api/companies/{cid}/brief").json()
    labels = [s["label"] for s in b["sections"][0]["statements"]]
    assert labels == ["unverified", "unavailable"]  # invented volume is withheld
    assert b["outreach_email"]["warnings"]  # numbers not in evidence are flagged
    assert analyst.get(f"/api/companies/{cid}").status_code in (404, 405)  # no stray route
    assert analyst.get(f"/api/companies/{cid}/brief").json()["company"] == "Acme Foods"
    report = admin.get("/api/admin/ai-usage").json()
    assert report["by_feature"][0]["feature"] == "company_brief" and report["month_spend_usd"] > 0


def test_news_keeps_only_retrieved_sources(db, analyst, ai_on, monkeypatch):
    from app.ai import claude as claude_mod

    def fake_news(**kw):
        return (
            [
                {"title": "EU labelling rule", "url": "https://gov.example/rule", "date": "2026-05"},
                {"title": "Made up", "url": "https://nowhere.example"},
            ],
            {"https://gov.example/rule": {"url": "https://gov.example/rule", "title": "Gov", "quotes": []}},
            {"model": "claude-haiku-5-5", "web_searches": 2},
        )

    monkeypatch.setattr(claude_mod, "research_news", fake_news)
    r = analyst.post("/api/markets/190531/news", json={}).json()
    assert [i["title"] for i in r["items"]] == ["EU labelling rule"] and r["rejected"] == 1
    assert analyst.post("/api/markets/190531/news", json={}).json()["cached"] is True
    assert analyst.get("/api/markets/190531/news").json()["items"][0]["url"] == "https://gov.example/rule"


def test_company_and_list_exports(db, analyst):
    cid = _company(db, relationship_type="buyer", contact={"email": "sales@acme.example"})
    lid = analyst.post("/api/prospect-lists", json={"name": "UK Q1"}).json()["id"]
    analyst.post(f"/api/prospect-lists/{lid}/companies/{cid}")
    body = analyst.get(f"/api/prospect-lists/{lid}/export").content.decode("utf-8-sig")
    assert "Acme Foods" in body and "sales@acme.example" in body
    assert "Acme Foods" in analyst.get("/api/companies/export?relationship=buyer").content.decode("utf-8-sig")


def test_name_check_handles_entities_brackets_and_suffixes():
    assert research.name_in_text("John & Pascalis", "<h1>John &amp; Pascalis</h1>") is True
    assert research.name_in_text("Katsouris Brothers Ltd (Cypressa)", "<p>Katsouris  Brothers supply</p>") is True
    assert research.name_in_text("Acme Foods Ltd", "<p>Other company</p>") is False


def test_login_rate_limit():
    from app.routers import auth as auth_router

    auth_router.FAILED.clear()
    c = TestClient(app)
    codes = [
        c.post("/api/auth/login", data={"username": "viewer@example.com", "password": "nope"}).status_code
        for _ in range(11)
    ]
    assert codes[:10] == [401] * 10 and codes[10] == 429
    auth_router.FAILED.clear()


def test_security_headers_and_api_404():
    c = TestClient(app)
    r = c.get("/api/health")
    assert r.headers["x-content-type-options"] == "nosniff" and r.headers["cache-control"] == "no-store"


def test_change_password(viewer):
    from tests.conftest import PASSWORD

    assert (
        viewer.post(
            "/api/auth/change-password", json={"current_password": "wrong", "new_password": "newpass99"}
        ).status_code
        == 400
    )
    assert (
        viewer.post(
            "/api/auth/change-password", json={"current_password": PASSWORD, "new_password": "short"}
        ).status_code
        == 422
    )
    assert (
        viewer.post(
            "/api/auth/change-password", json={"current_password": PASSWORD, "new_password": "newpass99"}
        ).status_code
        == 204
    )
    anon = TestClient(app)
    assert (
        anon.post("/api/auth/login", data={"username": "viewer@example.com", "password": "newpass99"}).status_code
        == 200
    )


def test_keyed_refresh_uses_three_requests_and_covers_every_market(db, monkeypatch):
    from app.config import get_settings
    from app.models import RefreshRun

    monkeypatch.setattr(get_settings(), "comtrade_subscription_key", "test-key")
    real_fetch = comtrade.fetch
    calls = []

    def fake_fetch(db_, **kw):
        calls.append((str(kw["period"]), kw["flow"], kw.get("reporter"), kw.get("partner")))
        if kw.get("partner") == 0:
            rows = [{**ROWS[0], "refYear": y, "flowCode": kw["flow"]} for y in map(int, str(kw["period"]).split(","))]
        else:  # all reporters x all partners for the latest two years
            rows = [{**ROWS[0], "refYear": 2025, "partnerCode": 484, "partnerISO": "MEX", "partnerDesc": "Mexico"}]
        return real_fetch(db_, caller=lambda params, key: rows, **kw)

    monkeypatch.setattr(comtrade, "fetch", fake_fetch)
    monkeypatch.setattr(trade, "_years", lambda: [2020, 2021, 2022, 2023, 2024, 2025])
    run_id = _start_run(db)
    trade.refresh_hs(run_id, "190531", sleep=lambda s: None)
    db.expire_all()
    assert db.get(RefreshRun, run_id).status == "succeeded"
    assert calls == [
        ("2020,2021,2022,2023,2024,2025", "M", None, 0),
        ("2020,2021,2022,2023,2024,2025", "X", None, 0),
        ("2024,2025", "M", None, None),
    ]
    assert comtrade.max_records() == comtrade.MAX_RECORDS_KEYED


def test_bulk_store_dedupes_rows_within_a_response(db):
    src = SourceRecord(adapter="un_comtrade", url="https://comtradeapi.un.org/test", reporting_period="2023")
    db.add(src)
    db.flush()
    dup = [{**ROWS[0], "primaryValue": 1.0}, {**ROWS[0], "primaryValue": 2.0}]  # same unique key twice
    assert trade._store(db, dup, "190531", src.id) == 1
    db.commit()
    assert db.scalar(select(TradeObservation.value_usd)) == 2.0


def test_heartbeat_decides_liveness_not_restarts(db, analyst):
    from datetime import UTC, datetime, timedelta

    from app.models import RefreshRun
    from app.services import jobs

    fresh = RefreshRun(adapter="un_comtrade", scope="HS 190531", status="running", heartbeat_at=datetime.now(UTC))
    stale = RefreshRun(
        adapter="un_comtrade",
        scope="HS 090111",
        status="queued",
        heartbeat_at=datetime.now(UTC) - timedelta(minutes=10),
    )
    db.add_all([fresh, stale])
    db.commit()
    assert jobs.is_live(fresh) and not jobs.is_live(stale)
    # another process starting up must not cancel a live job, only the abandoned one
    assert jobs.sweep(db) == 1
    db.expire_all()
    assert db.get(RefreshRun, fresh.id).status == "running" and db.get(RefreshRun, stale.id).status == "failed"
    # a live job blocks a duplicate refresh; the API reports a stale job as failed
    assert analyst.post("/api/markets/190531/refresh").json()["id"] == fresh.id
    jobs.beat([fresh.id])
    db.expire_all()
    assert jobs.is_live(db.get(RefreshRun, fresh.id))
